from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

from factoryos.core.services.change_log_service import log_change
from factoryos.extensions import db
from factoryos.modules.masterdata.shared.constants import (
    MATERIAL_BASE_UNITS,
    MATERIAL_DELIVERY_FORMS,
    MATERIAL_STATUSES,
    MATERIAL_TYPES,
    POLYMER_TYPES,
)

from ..models import Material
from .material_service import material_data_from_input
from .material_storage_service import create_material_folders


EXPECTED_HEADERS = [
    "material_no", "external_material_no", "name", "description",
    "material_type", "polymer_type", "manufacturer", "color", "color_no",
    "delivery_form", "base_unit", "density_g_cm3", "melt_flow_rate",
    "drying_temperature_c", "drying_time_h",
    "processing_temperature_min_c", "processing_temperature_max_c",
    "mold_temperature_min_c", "mold_temperature_max_c", "storage_location",
    "minimum_stock", "reorder_point", "material_status", "hazardous_material",
    "preferred_supplier_name", "supplier_material_no",
]

HEADER_LABELS = {
    "material_no": "Materialnummer",
    "external_material_no": "Externe Materialnummer",
    "name": "Genaue Bezeichnung / Handelsname",
    "description": "Beschreibung",
    "material_type": "Materialart",
    "polymer_type": "Kunststoffart",
    "manufacturer": "Hersteller",
    "color": "Farbe",
    "color_no": "Farbnummer",
    "delivery_form": "Lieferform",
    "base_unit": "Basiseinheit",
    "density_g_cm3": "Dichte (g/cm³)",
    "melt_flow_rate": "Schmelzindex / MFR",
    "drying_temperature_c": "Trocknungstemperatur (°C)",
    "drying_time_h": "Trocknungszeit (h)",
    "processing_temperature_min_c": "Verarbeitung min. (°C)",
    "processing_temperature_max_c": "Verarbeitung max. (°C)",
    "mold_temperature_min_c": "Werkzeugtemperatur min. (°C)",
    "mold_temperature_max_c": "Werkzeugtemperatur max. (°C)",
    "storage_location": "Lagerort",
    "minimum_stock": "Mindestbestand",
    "reorder_point": "Meldebestand",
    "material_status": "Status",
    "hazardous_material": "Gefahrstoff",
    "preferred_supplier_name": "Bevorzugter Lieferant",
    "supplier_material_no": "Materialnummer beim Lieferanten",
}

HEADER_COMMENTS = {
    "material_no": "Pflichtfeld und eindeutiger Schlüssel. Beim Import werden bestehende Datensätze damit aktualisiert.",
    "name": "Pflichtfeld: genaue Werkstoffbezeichnung oder Handelsname.",
    "material_type": "Pflichtfeld: plastic, operating oder auxiliary. Deutsche Bezeichnungen werden ebenfalls akzeptiert.",
    "polymer_type": "Bei Kunststoffmaterial Pflicht. Zulässige Werte stehen in der Dropdown-Liste.",
    "hazardous_material": "Zulässig: Ja/Nein, true/false, 1/0 oder x.",
    "preferred_supplier_name": "Vorbereitung für das spätere Lieferantenmodul; aktuell als Klartext gespeichert.",
}

TYPE_ALIASES = {
    "plastic": "plastic", "kunststoff": "plastic", "kunststoffmaterial": "plastic",
    "operating": "operating", "betriebsstoff": "operating", "betriebsstoffe": "operating",
    "auxiliary": "auxiliary", "hilfsstoff": "auxiliary", "hilfsstoffe": "auxiliary",
}
STATUS_ALIASES = {
    "aktiv": "aktiv", "active": "aktiv", "gesperrt": "gesperrt", "blocked": "gesperrt",
    "auslaufend": "auslaufend", "phase out": "auslaufend", "inaktiv": "inaktiv", "inactive": "inaktiv",
}
DELIVERY_FORM_ALIASES = {
    **{key.lower(): key for key in MATERIAL_DELIVERY_FORMS},
    **{label.lower(): key for key, label in MATERIAL_DELIVERY_FORMS.items()},
}
BASE_UNIT_ALIASES = {
    **{key.lower(): key for key in MATERIAL_BASE_UNITS},
    **{label.lower(): key for key, label in MATERIAL_BASE_UNITS.items()},
    "liter": "l", "stück": "st", "stueck": "st",
}


def _normalize(value):
    return str(value or "").replace("\xa0", " ").strip().lower()


def _normalize_polymer(value):
    normalized = _normalize(value)
    if not normalized:
        return None
    for key, label in POLYMER_TYPES.items():
        candidates = {
            key.lower(), label.lower(), label.split(" – ", 1)[0].lower(),
            key.lower().replace("_", "/"), key.lower().replace("_", "-"),
        }
        if normalized in candidates:
            return key
    return None


def _map_choice(value, aliases, field_label, default=None):
    if value in (None, ""):
        return default
    result = aliases.get(_normalize(value))
    if not result:
        raise ValueError(f"Unbekannter Wert für {field_label}: '{value}'")
    return result


def _row_dict(headers, row):
    return {
        header: row[index] if index < len(row) else None
        for index, header in enumerate(headers)
        if header
    }


def _canonical_row(raw):
    data = dict(raw)
    data["material_type"] = _map_choice(
        raw.get("material_type"), TYPE_ALIASES, "Materialart", "plastic"
    )
    data["material_status"] = _map_choice(
        raw.get("material_status"), STATUS_ALIASES, "Status", "aktiv"
    )
    data["base_unit"] = _map_choice(
        raw.get("base_unit"), BASE_UNIT_ALIASES, "Basiseinheit", "kg"
    )
    data["delivery_form"] = _map_choice(
        raw.get("delivery_form"), DELIVERY_FORM_ALIASES, "Lieferform"
    )

    polymer = _normalize_polymer(raw.get("polymer_type"))
    if raw.get("polymer_type") not in (None, "") and not polymer:
        raise ValueError(f"Unbekannte Kunststoffart: '{raw.get('polymer_type')}'")
    data["polymer_type"] = polymer
    return material_data_from_input(data)


def _find_import_sheet(workbook):
    for worksheet in workbook.worksheets:
        headers = [
            _normalize(cell.value) if cell.value else ""
            for cell in worksheet[1]
        ]
        if "material_no" in headers:
            return worksheet, headers
    raise ValueError("Keine Tabelle mit der Spalte 'material_no' gefunden.")


def import_materials_from_excel(file, user_id):
    workbook = load_workbook(file, data_only=True)
    worksheet, headers = _find_import_sheet(workbook)
    created = 0
    updated = 0
    errors = []

    for row_index, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        raw = _row_dict(headers, row)
        material_no = str(raw.get("material_no") or "").strip() or "-"

        try:
            values = _canonical_row(raw)
            with db.session.begin_nested():
                material = Material.query.filter_by(material_no=values["material_no"]).first()
                is_new = material is None
                if is_new:
                    material = Material(created_by_id=user_id)
                    db.session.add(material)
                    db.session.flush()

                changes = {
                    field: {"old": getattr(material, field, None), "new": value}
                    for field, value in values.items()
                    if getattr(material, field, None) != value
                }
                for field, value in values.items():
                    setattr(material, field, value)
                db.session.flush()

                log_change(
                    entity_type="material",
                    entity_id=material.id,
                    entity_name=material.material_no,
                    action="import" if is_new else "update",
                    changes=(
                        {"material_no": {"old": None, "new": material.material_no}}
                        if is_new else changes
                    ),
                    category="masterdata",
                )
                create_material_folders(material.material_no)

            if is_new:
                created += 1
            else:
                updated += 1
        except Exception as error:
            errors.append({"row": row_index, "material_no": material_no, "reason": str(error)})

    db.session.commit()
    return {"created": created, "updated": updated, "errors": errors}


def _style_data_sheet(ws, row_count):
    dark = PatternFill("solid", fgColor="1F2937")
    white_bold = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXPECTED_HEADERS))}{max(row_count, 2)}"
    ws.row_dimensions[1].height = 32
    for index, header in enumerate(EXPECTED_HEADERS, start=1):
        cell = ws.cell(1, index)
        cell.fill = dark
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        cell.comment = Comment(
            HEADER_COMMENTS.get(header, HEADER_LABELS[header]),
            "FactoryOS",
        )
        width = max(14, min(34, len(HEADER_LABELS[header]) + 4))
        ws.column_dimensions[get_column_letter(index)].width = width

    for row in ws.iter_rows(min_row=2, max_row=max(row_count, 2), max_col=len(EXPECTED_HEADERS)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    numeric_formats = {
        "density_g_cm3": "0.000",
        "drying_temperature_c": "0.0",
        "drying_time_h": "0.0",
        "processing_temperature_min_c": "0.0",
        "processing_temperature_max_c": "0.0",
        "mold_temperature_min_c": "0.0",
        "mold_temperature_max_c": "0.0",
        "minimum_stock": "0.000",
        "reorder_point": "0.000",
    }
    for header, number_format in numeric_formats.items():
        column = EXPECTED_HEADERS.index(header) + 1
        for row_index in range(2, max(row_count, 2) + 1):
            ws.cell(row_index, column).number_format = number_format

    required_fill = PatternFill("solid", fgColor="FFF2CC")
    for header in ("material_no", "name", "material_type"):
        ws.cell(1, EXPECTED_HEADERS.index(header) + 1).fill = required_fill
        ws.cell(1, EXPECTED_HEADERS.index(header) + 1).font = Font(color="111827", bold=True)

    material_type_col = get_column_letter(EXPECTED_HEADERS.index("material_type") + 1)
    polymer_col = get_column_letter(EXPECTED_HEADERS.index("polymer_type") + 1)
    ws.conditional_formatting.add(
        f"{polymer_col}2:{polymer_col}1001",
        FormulaRule(
            formula=[f'${material_type_col}2="plastic"'],
            fill=PatternFill("solid", fgColor="FFF2CC"),
        ),
    )


def _add_validation(ws, header, defined_name):
    column = get_column_letter(EXPECTED_HEADERS.index(header) + 1)
    validation = DataValidation(
        type="list",
        formula1=f"={defined_name}",
        allow_blank=header not in {"material_type", "material_status", "base_unit"},
    )
    validation.error = "Bitte einen Wert aus der Liste auswählen."
    validation.errorTitle = "Ungültiger Wert"
    validation.prompt = "Wert aus der Dropdown-Liste auswählen."
    validation.promptTitle = HEADER_LABELS[header]
    validation.showErrorMessage = True
    validation.showInputMessage = True
    ws.add_data_validation(validation)
    validation.add(f"{column}2:{column}1001")


def _material_values(material):
    values = {}
    for header in EXPECTED_HEADERS:
        if header == "hazardous_material":
            value = "Ja" if material.hazardous_material else "Nein"
        else:
            value = getattr(material, header, None)
        values[header] = "" if value is None else value
    return values


def build_material_workbook(materials=None, template=False):
    materials = list(materials or [])
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Materialien"
    ws.append(EXPECTED_HEADERS)

    if not template:
        for material in materials:
            values = _material_values(material)
            ws.append([values[header] for header in EXPECTED_HEADERS])

    _style_data_sheet(ws, max(len(materials) + 1, 2))

    overview = workbook.create_sheet("Übersicht")
    overview.append([HEADER_LABELS[header] for header in EXPECTED_HEADERS])
    if not template:
        for material in materials:
            values = _material_values(material)
            values["material_type"] = MATERIAL_TYPES.get(material.material_type, material.material_type)
            values["polymer_type"] = POLYMER_TYPES.get(material.polymer_type, material.polymer_type or "")
            values["delivery_form"] = MATERIAL_DELIVERY_FORMS.get(material.delivery_form, material.delivery_form or "")
            values["base_unit"] = MATERIAL_BASE_UNITS.get(material.base_unit, material.base_unit or "")
            values["material_status"] = MATERIAL_STATUSES.get(material.material_status, material.material_status)
            overview.append([values[header] for header in EXPECTED_HEADERS])
    _style_data_sheet(overview, max(len(materials) + 1, 2))

    lists = workbook.create_sheet("Auswahllisten")
    choices = [
        ("Materialart", list(MATERIAL_TYPES)),
        ("Kunststoffart", list(POLYMER_TYPES)),
        ("Status", list(MATERIAL_STATUSES)),
        ("Basiseinheit", list(MATERIAL_BASE_UNITS)),
        ("Lieferform", list(MATERIAL_DELIVERY_FORMS)),
        ("Gefahrstoff", ["Ja", "Nein"]),
    ]
    for column_index, (title, values) in enumerate(choices, start=1):
        lists.cell(1, column_index, title)
        for row_index, value in enumerate(values, start=2):
            lists.cell(row_index, column_index, value)

    for header, list_column, count, defined_name in (
        ("material_type", "A", len(MATERIAL_TYPES), "Materialarten"),
        ("polymer_type", "B", len(POLYMER_TYPES), "Kunststoffarten"),
        ("material_status", "C", len(MATERIAL_STATUSES), "Materialstatus"),
        ("base_unit", "D", len(MATERIAL_BASE_UNITS), "Basiseinheiten"),
        ("delivery_form", "E", len(MATERIAL_DELIVERY_FORMS), "Lieferformen"),
        ("hazardous_material", "F", 2, "Gefahrstoffauswahl"),
    ):
        workbook.defined_names.add(DefinedName(
            defined_name,
            attr_text=f"'Auswahllisten'!${list_column}$2:${list_column}${count + 1}",
        ))
        _add_validation(ws, header, defined_name)

    lists.sheet_state = "hidden"
    workbook.active = 0
    return workbook


def material_workbook_bytes(materials=None, template=False):
    output = BytesIO()
    build_material_workbook(materials, template=template).save(output)
    output.seek(0)
    return output


def export_materials_to_excel():
    materials = Material.query.order_by(Material.material_no.asc()).all()
    return material_workbook_bytes(materials)
